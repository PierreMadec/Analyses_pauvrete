import csv
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = "data"
OUT = "insee_référence_figures.xlsx"

FONT = "Arial"
TITLE_FONT = Font(name=FONT, size=12, bold=True)
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="595959")
HEADER_FILL_COLOR = "1F4E78"
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)

# (numero_fichier, onglet, titre, unite, lecture, champ, source)
FIGURES = [
    ("01_emploi_pauvrete", "Figure 1",
     "Taux d'emploi et taux de pauvreté monétaire, 2005-2023", "%",
     "En 2023, le taux d'emploi des 18-64 ans atteint son plus haut niveau depuis 2005, tandis que le taux de pauvreté repart à la hausse.",
     "France métropolitaine, 18-64 ans.", "Insee, ERFS."),
    ("02_relatif_ancre", "Figure 2",
     "Taux de pauvreté des personnes en emploi selon le seuil retenu, 2005-2023", "%",
     "Au seuil ancré (pouvoir d'achat constant), le taux de pauvreté laborieuse recule de 8,8 % en 2005 à 6,4 % en 2023.",
     "France métropolitaine, personnes de référence en emploi, 18-64 ans.", "Insee, ERFS."),
    ("03_taux_par_config", "Figure 3",
     "Taux de pauvreté laborieuse selon la configuration du ménage, 2005-2023", "%",
     "En 2023, le taux de pauvreté des couples biactifs sans enfant est quasi nul.",
     "France métropolitaine, personnes de référence en emploi, 18-64 ans.", "Insee, ERFS."),
    ("04_diagnostic_causes", "Figure 4",
     "Causes proximales de la pauvreté laborieuse par configuration de ménage, 2021-2023", "%",
     "Chez les personnes seules pauvres, 70 % ont un salaire en équivalent temps plein inférieur au Smic annuel.",
     "France métropolitaine, personnes de référence en emploi pauvres, 18-64 ans.", "Insee, ERFS."),
    ("05_distribution_salaire_horaire", "Figure 5",
     "Distribution du salaire horaire rapporté au Smic, travailleurs pauvres et non pauvres, 2013-2023", "Densité",
     "41 % des travailleurs pauvres ont un salaire horaire supérieur au Smic net.",
     "France métropolitaine, personnes de référence en emploi, 18-64 ans, avec quotité de temps partiel connue.", "Insee, ERFS."),
    ("06_contrefactuels", "Figure 6",
     "Travailleurs pauvres sortis du seuil selon le levier simulé (maquette SoFi), par configuration, 2021-2023", "%",
     "Un passage à temps plein au Smic sort 67 % des personnes seules pauvres du seuil de pauvreté.",
     "France métropolitaine, personnes de référence en emploi pauvres, 18-64 ans.",
     "Insee, ERFS ; Pucci M., SoFi (2022), calculs de l'auteur."),
    ("07_oaxaca_composition_coefficients", "Figure 7",
     "Décomposition de Oaxaca-Blinder du taux de pauvreté parmi les travailleurs, 2010-2012 / 2021-2023", "Points de %",
     "L'effet de composition contribue pour +0,50 point à la variation du taux de pauvreté laborieuse.",
     "France métropolitaine, personnes de référence en emploi, 18-64 ans.", "Insee, ERFS, calculs de l'auteur."),
    ("08_cohortes", "Figure 8",
     "Taux de pauvreté laborieuse à 25-34 ans selon la cohorte de naissance, 2005-2023", "%",
     "La cohorte née dans les années 1990 affiche un taux de pauvreté laborieuse à 25-34 ans supérieur à 10 %.",
     "France métropolitaine, personnes de référence en emploi, 25-34 ans.", "Insee, ERFS."),
    ("09_effet_amortisseur", "Figure 9",
     "Effet amortisseur de l'enveloppe de soutien aux bas revenus d'activité, 2005-2023", "%",
     "En 2019, l'enveloppe de soutien à l'activité réduit le taux de pauvreté laborieuse de 2,0 points.",
     "France métropolitaine, personnes de référence en emploi, 18-64 ans.", "Insee, ERFS, calculs de l'auteur."),
    ("10_decompo_transferts", "Figure 10",
     "Effet amortisseur de chaque poste de transferts par configuration de ménage, 2021-2023", "Points de %",
     "Chez les familles monoparentales monoactives, les prestations familiales réduisent le taux de pauvreté de 10,2 points.",
     "France métropolitaine, personnes de référence en emploi, 18-64 ans.", "Insee, ERFS, calculs de l'auteur."),
]


def read_csv(path):
    with open(path, encoding="utf-8") as f:
        rows = list(csv.reader(f))
    return rows[0], rows[1:]


def round_numeric(val):
    try:
        f = float(val)
        return round(f, 2)
    except (ValueError, TypeError):
        return val


def build_sheet(ws, titre, unite, lecture, champ, source, header, rows):
    ws.sheet_view.showGridLines = False
    ncols = len(header)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 2))
    c = ws.cell(row=1, column=1, value=titre)
    c.font = TITLE_FONT
    ws.row_dimensions[1].height = 22

    ws.cell(row=2, column=1, value="Unité :").font = Font(name=FONT, size=9, bold=True)
    ws.cell(row=2, column=2, value=unite).font = Font(name=FONT, size=9)

    header_row = 4
    for j, colname in enumerate(header, start=1):
        cell = ws.cell(row=header_row, column=j, value=colname)
        cell.font = HEADER_FONT
        cell.fill = __import__("openpyxl").styles.PatternFill("solid", fgColor=HEADER_FILL_COLOR)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for i, row in enumerate(rows, start=header_row + 1):
        for j, val in enumerate(row, start=1):
            v = round_numeric(val) if j > 1 else val
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = BODY_FONT
            cell.border = BORDER
            if j > 1:
                cell.alignment = Alignment(horizontal="right")

    last_row = header_row + len(rows)
    for j, colname in enumerate(header, start=1):
        maxlen = max([len(str(colname))] + [len(str(r[j - 1])) for r in rows]) if rows else len(str(colname))
        ws.column_dimensions[get_column_letter(j)].width = min(max(maxlen + 2, 10), 40)

    note_row = last_row + 2
    for label, text in [("Lecture :", lecture), ("Champ :", champ), ("Source :", source)]:
        ws.cell(row=note_row, column=1, value=label).font = Font(name=FONT, size=9, bold=True, italic=True, color="595959")
        cell = ws.cell(row=note_row, column=2, value=text)
        cell.font = NOTE_FONT
        ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=max(ncols, 4))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        note_row += 1


def main():
    wb = Workbook()
    wb.remove(wb.active)
    for fname, sheet_name, titre, unite, lecture, champ, source in FIGURES:
        header, rows = read_csv(f"{DATA_DIR}/{fname}.csv")
        ws = wb.create_sheet(sheet_name)
        build_sheet(ws, titre, unite, lecture, champ, source, header, rows)
    wb.save(OUT)
    print("Saved", OUT)


if __name__ == "__main__":
    main()
